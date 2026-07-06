"""Build a corrected Stonesaga Codex dataset from the PDF and errata DOCX.

The extractor is intentionally conservative. It derives records from the PDF
text layer, applies only clear errata patterns, and logs the rest for manual
review instead of guessing.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

import fitz
import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PDF = ROOT / "assets" / "Stonesaga Codex v1.pdf"
DEFAULT_DOCX = ROOT / "assets" / "Stonesaga Errata & FAQ v1.3.1.docx"
DEFAULT_DATA_DIR = ROOT / "data"

ERRATA_VERSION_RE = re.compile(r"\b(Added|Corrected|Updated)\s+v?([0-9][0-9A-Za-z.\-]*)", re.I)
PAGE_REF_RE = re.compile(r"\bOn\s+P\.\s*([0-9]+(?:\s*-\s*[0-9]+)?)", re.I)
OUTCOME_ID_RE = re.compile(r"^\d{3}$")
CODE_RE = re.compile(r"^(?:[A-Z]{1,3}\d{2}[a-z]?|[YRBPSON]\d{4})$")
RECIPE_CODE_RE = re.compile(r"^\d{4}$")
OPTION_RE = re.compile(r"([A-Z])\)")


@dataclass(frozen=True)
class TocSection:
    section: str
    subsection: str
    start: int
    end: int


TOC: list[TocSection] = [
    TocSection("OUTCOMES", "", 3, 111),
    TocSection("OMENS", "CLOUD", 113, 118),
    TocSection("OMENS", "COMET", 119, 122),
    TocSection("OMENS", "MOON", 123, 128),
    TocSection("OMENS", "STAR", 129, 132),
    TocSection("OMENS", "SUN", 133, 136),
    TocSection("RECIPES", "YELLOW", 137, 150),
    TocSection("RECIPES", "RED", 151, 164),
    TocSection("RECIPES", "BLUE", 165, 174),
    TocSection("RECIPES", "PURPLE", 175, 176),
    TocSection("RECIPES", "SILVER", 177, 178),
    TocSection("RECIPES", "ORANGE", 179, 181),
    TocSection("MINING CHART", "", 183, 184),
    TocSection("MYSTERY CHART", "", 185, 188),
    TocSection("CREATING NAMES", "", 189, 190),
    TocSection("GRINDING RESULTS", "", 191, 193),
    TocSection("EPOCH 2", "", 195, 202),
    TocSection("EPOCH 3", "", 203, 208),
    TocSection("OVERLAY REFERENCE", "", 209, 212),
    TocSection("ADVANCED REGION RULES", "", 213, 218),
    # Corrected by Codex errata v1.3.0: the printed TOC swaps these two ranges.
    TocSection("HEX KEY", "", 219, 223),
    TocSection("MARK LIST", "", 225, 227),
]

TOC_ORDER = {(toc.section, toc.subsection): index for index, toc in enumerate(TOC)}


SECTION_BY_PAGE: dict[int, TocSection] = {}
for toc in TOC:
    for page in range(toc.start, toc.end + 1):
        SECTION_BY_PAGE[page] = toc


FIELDNAMES = [
    "stable_key",
    "section",
    "subsection",
    "page_start",
    "page_end",
    "source_id_or_name",
    "original_text",
    "corrected_text",
    "errata_applied",
    "errata_source",
    "errata_version_or_date",
    "change_summary",
    "notes",
]


@dataclass
class Record:
    stable_key: str
    section: str
    subsection: str
    page_start: int
    page_end: int
    source_id_or_name: str
    original_text: str
    corrected_text: str
    errata_applied: bool = False
    errata_source: str = ""
    errata_version_or_date: str = ""
    change_summary: str = ""
    notes: str = ""

    def as_dict(self) -> dict[str, object]:
        return {name: getattr(self, name) for name in FIELDNAMES}


@dataclass
class Erratum:
    target: str
    codex_group: str
    page_ref: str
    body: str
    version: str
    action: str
    raw_paragraphs: list[str] = field(default_factory=list)


def slugify(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-") or "unknown"


def normalize_text(text: str) -> str:
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2026": "...",
        "\uf0b7": "-",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_printed_pages(pdf_path: Path) -> dict[int, str]:
    doc = fitz.open(pdf_path)
    pages: dict[int, str] = {}

    for physical_index, physical_page in enumerate(doc):
        rect = physical_page.rect
        clips = [
            ("left", fitz.Rect(0, 0, rect.width / 2, rect.height)),
            ("right", fitz.Rect(rect.width / 2, 0, rect.width, rect.height)),
        ]
        if rect.width < 900:
            clips = [("single", rect)]

        for side, clip in clips:
            text = normalize_text(physical_page.get_text("text", clip=clip))
            if not text:
                continue
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            page_no = None
            if lines:
                first_line_digits = lines[0].replace(" ", "")
                if re.fullmatch(r"\d{1,3}", first_line_digits):
                    page_no = int(first_line_digits)
            if page_no is None and physical_index >= 2 and side in {"left", "right"}:
                # The text layer does not always emit the printed page number,
                # but the post-TOC PDF is a stable two-page spread: PDF index 2
                # contains printed pages 4 and 5, index 3 contains 6 and 7, etc.
                page_no = physical_index * 2 + (0 if side == "left" else 1)
            if page_no is None or page_no not in SECTION_BY_PAGE:
                continue
            pages[page_no] = clean_page_text(page_no, lines)

    return pages


def clean_page_text(page_no: int, lines: list[str]) -> str:
    cleaned: list[str] = []
    toc = SECTION_BY_PAGE[page_no]
    spaced_section = " ".join(toc.section)
    removed_page_number = False

    for line in lines:
        compact = re.sub(r"\s+", "", line).upper()
        if not removed_page_number and line.replace(" ", "") == str(page_no):
            removed_page_number = True
            continue
        if compact in {"ID", "ENTRY", "IDENTRY"}:
            continue
        if compact in {
            "YELLOWRECIPES",
            "REDRECIPES",
            "BLUERECIPES",
            "PURPLERECIPES",
            "SILVERRECIPES",
            "ORANGERECIPES",
        }:
            continue
        if line.upper() in {toc.section, toc.subsection, "OUTCOMES", "OMENS", "RECIPES"}:
            continue
        if line.upper() == spaced_section:
            continue
        if compact in {"OUTCOMES", "OMENS", "RECIPES"}:
            continue
        cleaned.append(line)

    return normalize_text("\n".join(cleaned))


def section_for_page(page_no: int) -> TocSection:
    return SECTION_BY_PAGE[page_no]


def new_record(section: TocSection, source_id: str, page_no: int, text: str) -> Record:
    prefix = section.section
    if section.subsection:
        prefix = f"{prefix}-{section.subsection}"
    stable_key = slugify(f"{prefix}-{source_id}")
    return Record(
        stable_key=stable_key,
        section=section.section,
        subsection=section.subsection,
        page_start=page_no,
        page_end=page_no,
        source_id_or_name=source_id,
        original_text=normalize_text(text),
        corrected_text=normalize_text(text),
    )


def parse_outcomes(pages: dict[int, str]) -> list[Record]:
    records: list[Record] = []
    by_id: dict[str, Record] = {}
    current_id: str | None = None
    current_lines: list[str] = []
    current_page = 0

    def flush() -> None:
        nonlocal current_id, current_lines, current_page
        if not current_id:
            return
        text = normalize_text("\n".join(line for line in current_lines if line != "(ctnd.)"))
        if current_id in by_id:
            record = by_id[current_id]
            record.original_text = normalize_text(record.original_text + "\n\n" + text)
            record.corrected_text = record.original_text
            record.page_end = max(record.page_end, current_page)
        else:
            section = section_for_page(current_page)
            record = new_record(section, current_id, current_page, text)
            records.append(record)
            by_id[current_id] = record
        current_id = None
        current_lines = []

    for page_no in sorted(p for p in pages if 3 <= p <= 111):
        for line in pages[page_no].splitlines():
            stripped = line.strip()
            id_match = re.match(r"^(\d{3})(?:\s+(.*))?$", stripped)
            if id_match:
                flush()
                current_id = id_match.group(1)
                current_lines = [id_match.group(2)] if id_match.group(2) else []
                current_page = page_no
                continue
            if current_id:
                current_lines.append(stripped)
                current_page = page_no
        flush()

    return records


def parse_code_section(page_no: int, text: str) -> list[Record]:
    section = section_for_page(page_no)
    if not text.strip():
        return []
    records: list[Record] = []
    current_code: str | None = None
    current_lines: list[str] = []
    recipe_prefix_by_subsection = {
        "YELLOW": "Y",
        "RED": "R",
        "BLUE": "B",
        "PURPLE": "P",
        "SILVER": "S",
        "ORANGE": "O",
    }

    def flush() -> None:
        nonlocal current_code, current_lines
        if current_code:
            records.append(new_record(section, current_code, page_no, "\n".join(current_lines)))
        current_code = None
        current_lines = []

    for line in text.splitlines():
        stripped = line.strip()
        recipe_code = section.section == "RECIPES" and RECIPE_CODE_RE.fullmatch(stripped)
        omen_code = section.section == "OMENS" and CODE_RE.fullmatch(stripped)
        if recipe_code or omen_code:
            flush()
            if recipe_code:
                current_code = f"{recipe_prefix_by_subsection[section.subsection]}{stripped}"
            else:
                current_code = stripped
            current_lines = []
        elif current_code:
            current_lines.append(stripped)

    flush()

    if records:
        return records

    source = f"{section.section} p.{page_no}"
    return [new_record(section, source, page_no, text)]


def parse_non_outcome_records(pages: dict[int, str]) -> list[Record]:
    records: list[Record] = []
    merged_page_sections = {
        "MINING CHART",
        "MYSTERY CHART",
        "CREATING NAMES",
        "GRINDING RESULTS",
        "EPOCH 2",
        "EPOCH 3",
        "OVERLAY REFERENCE",
        "ADVANCED REGION RULES",
        "MARK LIST",
        "HEX KEY",
    }
    section_buffers: dict[str, list[tuple[int, str]]] = defaultdict(list)

    for page_no in sorted(p for p in pages if p >= 113):
        section = section_for_page(page_no)
        if section.section in {"OMENS", "RECIPES"}:
            records.extend(parse_code_section(page_no, pages[page_no]))
        elif section.section in merged_page_sections:
            section_buffers[section.section].append((page_no, pages[page_no]))

    for section_name, chunks in section_buffers.items():
        first_page = min(page for page, _ in chunks)
        section = section_for_page(first_page)
        text = "\n\n".join(f"[p. {page}]\n{text}" for page, text in chunks)
        record = new_record(section, section_name.title(), first_page, text)
        record.page_end = max(page for page, _ in chunks)
        records.append(record)

    return records


def extract_records(pdf_path: Path) -> list[Record]:
    pages = extract_printed_pages(pdf_path)
    records = parse_outcomes(pages)
    records.extend(parse_non_outcome_records(pages))
    return merge_duplicate_records(records)


def merge_duplicate_records(records: list[Record]) -> list[Record]:
    merged: dict[str, Record] = {}
    ordered: list[Record] = []
    for record in records:
        existing = merged.get(record.stable_key)
        if not existing:
            merged[record.stable_key] = record
            ordered.append(record)
            continue
        existing.original_text = normalize_text(existing.original_text + "\n\n" + record.original_text)
        existing.corrected_text = existing.original_text
        existing.page_start = min(existing.page_start, record.page_start)
        existing.page_end = max(existing.page_end, record.page_end)
        existing.notes = append_field(existing.notes, "Merged duplicate source ID across page boundary.")
    return ordered


def iter_docx_paragraphs(docx_path: Path) -> Iterable[tuple[str, str]]:
    doc = Document(docx_path)
    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        if text:
            yield paragraph.style.name, text


def parse_codex_errata(docx_path: Path) -> list[Erratum]:
    paragraphs = list(iter_docx_paragraphs(docx_path))
    start = next(i for i, (_, text) in enumerate(paragraphs) if text == "Codex Corrections & Clarifications")
    end = next(i for i, (_, text) in enumerate(paragraphs[start + 1 :], start + 1) if text.startswith("Frequently Asked Questions"))
    codex = paragraphs[start:end]

    errata: list[Erratum] = []
    current_group = ""
    current_target = ""
    bucket: list[str] = []

    def flush_bucket() -> None:
        nonlocal bucket
        if not current_target or not bucket:
            bucket = []
            return
        page_ref = ""
        version = ""
        action = ""
        body_parts: list[str] = []
        for item in bucket:
            page_match = PAGE_REF_RE.search(item)
            version_match = ERRATA_VERSION_RE.search(item)
            if page_match:
                page_ref = page_match.group(1).replace(" ", "")
                continue
            if version_match:
                action = version_match.group(1).title()
                version = "v" + version_match.group(2)
                if body_parts:
                    errata.append(Erratum(current_target, current_group, page_ref, normalize_text("\n".join(body_parts)), version, action, bucket[:]))
                    body_parts = []
                continue
            body_parts.append(item)
        if body_parts:
            errata.append(Erratum(current_target, current_group, page_ref, normalize_text("\n".join(body_parts)), version, action, bucket[:]))
        bucket = []

    for style, text in codex:
        if style == "Heading 2":
            flush_bucket()
            current_group = text
            current_target = ""
        elif style == "Heading 3":
            flush_bucket()
            current_target = text
        elif current_target:
            bucket.append(text)
    flush_bucket()

    return [err for err in errata if err.body]


def target_keys(erratum: Erratum) -> list[str]:
    target = erratum.target.strip()
    keys: list[str] = []

    outcome = re.search(r"Outcome\s+([0-9]{3})(?:/([0-9]{3}))?", target, re.I)
    if outcome:
        keys.append(slugify(f"OUTCOMES-{outcome.group(1)}"))
        if outcome.group(2):
            keys.append(slugify(f"OUTCOMES-{outcome.group(2)}"))
        return keys

    recipe = re.fullmatch(r"([YRBPSON]\d{4})", target)
    if recipe:
        color_by_prefix = {
            "Y": "YELLOW",
            "R": "RED",
            "B": "BLUE",
            "P": "PURPLE",
            "S": "SILVER",
            "O": "ORANGE",
        }
        color = color_by_prefix.get(recipe.group(1)[0], "")
        keys.append(slugify(f"RECIPES-{color}-{recipe.group(1)}"))
        return keys

    omen = re.fullmatch(r"(Cloud|Comet|Moon|Star|Sun)\s+([A-Z]{1,3}\d{2}[a-z]?)", target, re.I)
    if omen:
        keys.append(slugify(f"OMENS-{omen.group(1).upper()}-{omen.group(2)}"))
        return keys

    if target in {"Epoch 2", "Epoch 3", "Overlay Reference", "Hex Key"}:
        keys.append(slugify(f"{target}-{target}"))
        keys.append(slugify(f"{target}-{target.title()}"))
    elif erratum.codex_group == "Special Regions":
        keys.append(slugify("ADVANCED REGION RULES-Advanced Region Rules"))
    elif erratum.codex_group:
        keys.append(slugify(f"{erratum.codex_group}-{target}"))

    return keys


def quoted_strings(text: str) -> list[str]:
    return re.findall(r'"([^"]+)"', text, flags=re.S)


def replace_option_block(text: str, option_letter: str, replacement: str) -> tuple[str, bool]:
    marker = f"{option_letter})"
    start = text.find(marker)
    if start < 0:
        return text, False
    match = re.search(r"\n[A-Z]\)", text[start + len(marker) :])
    end = start + len(marker) + match.start() if match else len(text)
    return text[:start] + replacement.strip() + text[end:], True


def apply_erratum_to_text(text: str, erratum: Erratum) -> tuple[str, bool, str]:
    body = normalize_text(erratum.body)
    quotes = quoted_strings(body)

    instead = re.search(r'text\s+"([^"]+)"\s+should instead (?:read|refer to)\s+"([^"]+)"', body, re.I | re.S)
    if instead and instead.group(1) in text:
        return text.replace(instead.group(1), instead.group(2)), True, f'Replaced "{instead.group(1)}" with "{instead.group(2)}".'

    refers = re.search(r'refers? to\s+"([^"]+)".*should instead (?:read|refer to)\s+"([^"]+)"', body, re.I | re.S)
    if refers and refers.group(1) in text:
        return text.replace(refers.group(1), refers.group(2)), True, f'Replaced "{refers.group(1)}" with "{refers.group(2)}".'

    label = re.search(r'entry labelled\s+"([^"]+)"\s+should instead be labelled\s+"([^"]+)"', body, re.I | re.S)
    if label and label.group(1) in text:
        return text.replace(label.group(1), label.group(2), 1), True, f'Relabelled {label.group(1)} as {label.group(2)}.'

    not_begin = re.search(r'should not begin with the sentence\s+"([^"]+)".*begin with the next sentence', body, re.I | re.S)
    if not_begin and text.startswith(not_begin.group(1)):
        return text[len(not_begin.group(1)) :].lstrip(), True, f'Removed opening sentence "{not_begin.group(1)}".'

    whole_read = re.search(r'(?:This outcome|This entry|This section).*should read:\s*(.+)', body, re.I | re.S)
    if whole_read:
        replacement = whole_read.group(1).strip()
        if replacement:
            return replacement, True, "Replaced full entry text with errata text."

    flavor_read = re.search(r"flavor text should read:?\s+\"([^\"]+)\"", body, re.I | re.S)
    if flavor_read:
        replacement = flavor_read.group(1).strip()
        mechanics = re.search(r"\n(You may immediately recombine|Discard|Gain|Return|If |When )", text)
        if mechanics:
            updated = replacement + text[mechanics.start() :]
            return updated, True, "Replaced recipe flavor text."
        return replacement, True, "Replaced recipe flavor text."

    option_read = re.search(r'The\s+"?([A-Z])\)?"?\s+option should read:?\s+"([^"]+)"', body, re.I | re.S)
    if option_read:
        updated, ok = replace_option_block(text, option_read.group(1).upper(), option_read.group(2))
        if ok:
            return updated, True, f"Replaced option {option_read.group(1).upper()})."

    if re.search(r"(sentence|step|option|mark|entry).*(should exist|should be added|add the following)", body, re.I) and quotes:
        addition = quotes[-1].strip()
        if addition and addition not in text:
            if re.search(r"begin|start", body, re.I):
                return addition + "\n" + text, True, "Prepended errata sentence/text."
            return text.rstrip() + "\n" + addition, True, "Appended errata sentence/text."
        if addition in text:
            return text, True, "Errata text already present."

    first_sentence = re.search(r'first sentence should read\s+"([^"]+)"', body, re.I | re.S)
    if first_sentence:
        replacement = first_sentence.group(1)
        existing = re.match(r"(.+?[.!?])(\s|$)", text, re.S)
        if existing:
            return replacement + text[existing.end(1) :], True, "Replaced first sentence."

    return text, False, "Instruction requires manual review or source text could not be matched safely."


def make_added_record(erratum: Erratum) -> Record:
    body = normalize_text(erratum.body)
    quotes = quoted_strings(body)
    text = "\n".join(quotes) if quotes else body
    page_start = int(erratum.page_ref.split("-")[0]) if erratum.page_ref and erratum.page_ref[0].isdigit() else 0
    toc = SECTION_BY_PAGE.get(page_start)
    omen = re.fullmatch(r"(Cloud|Comet|Moon|Star|Sun)\s+([A-Z]{1,3}\d{2}[a-z]?)", erratum.target, re.I)
    if toc is None and omen:
        subsection_start = {
            "CLOUD": 113,
            "COMET": 119,
            "MOON": 123,
            "STAR": 129,
            "SUN": 133,
        }[omen.group(1).upper()]
        toc = section_for_page(subsection_start)
        page_start = subsection_start
    if toc is None:
        toc = TocSection(erratum.codex_group.upper() or "CODEX", "", page_start, page_start)
    source = re.sub(r"^(Cloud|Comet|Moon|Star|Sun)\s+", "", erratum.target, flags=re.I)
    record = new_record(toc, source, page_start, text)
    record.original_text = ""
    record.corrected_text = text
    record.errata_applied = True
    record.errata_source = erratum.target
    record.errata_version_or_date = erratum.version
    record.change_summary = "Created missing Codex entry from errata."
    record.notes = "No matching PDF record was found; entry is supplied by the Codex errata."
    return record


def apply_errata(records: list[Record], errata: list[Erratum]) -> list[dict[str, object]]:
    by_key = {record.stable_key: record for record in records}
    log: list[dict[str, object]] = []

    for erratum in errata:
        if erratum.target == "Table of Contents":
            log.append(log_row(erratum, None, "applied", True, "Applied by corrected TOC range configuration."))
            continue
        relabelled = apply_relabel_erratum(erratum, by_key)
        if relabelled:
            record, summary = relabelled
            log.append(log_row(erratum, record, "applied", True, summary))
            continue
        keys = target_keys(erratum)
        matched = [by_key[key] for key in keys if key in by_key]
        if not matched and "An entry should exist" in erratum.body:
            record = make_added_record(erratum)
            records.append(record)
            by_key[record.stable_key] = record
            log.append(log_row(erratum, record, "created", True, record.change_summary))
            continue

        if not matched:
            log.append(log_row(erratum, None, "manual_review", False, "No matching record found."))
            continue

        for record in matched:
            updated, applied, summary = apply_erratum_to_text(record.corrected_text, erratum)
            if applied:
                record.corrected_text = normalize_text(updated)
                record.errata_applied = True
                record.errata_source = append_field(record.errata_source, erratum.target)
                record.errata_version_or_date = append_field(record.errata_version_or_date, erratum.version)
                record.change_summary = append_field(record.change_summary, summary)
            else:
                record.notes = append_field(record.notes, f"Manual review needed for {erratum.target}: {summary}")
            log.append(log_row(erratum, record, "applied" if applied else "manual_review", applied, summary))

    return log


def apply_relabel_erratum(erratum: Erratum, by_key: dict[str, Record]) -> tuple[Record, str] | None:
    label = re.search(r'entry labelled\s+"([^"]+)"\s+should instead be labelled\s+"([^"]+)"', erratum.body, re.I | re.S)
    if not label:
        return None

    old_label = label.group(1).strip()
    new_label = label.group(2).strip()
    target = erratum.target.strip()
    old_keys: list[str] = []
    new_key = ""

    omen = re.fullmatch(r"(Cloud|Comet|Moon|Star|Sun)\s+([A-Z]{1,3}\d{2}[a-z]?)", target, re.I)
    if omen:
        subsection = omen.group(1).upper()
        old_keys.append(slugify(f"OMENS-{subsection}-{old_label}"))
        new_key = slugify(f"OMENS-{subsection}-{new_label}")
    elif re.fullmatch(r"[YRBPSON]\d{4}", target):
        color_by_prefix = {
            "Y": "YELLOW",
            "R": "RED",
            "B": "BLUE",
            "P": "PURPLE",
            "S": "SILVER",
            "O": "ORANGE",
        }
        color = color_by_prefix.get(target[0], "")
        old_keys.append(slugify(f"RECIPES-{color}-{old_label}"))
        new_key = slugify(f"RECIPES-{color}-{new_label}")

    for old_key in old_keys:
        record = by_key.get(old_key)
        if not record:
            continue
        if new_key in by_key and by_key[new_key] is not record:
            return None
        del by_key[old_key]
        record.stable_key = new_key
        record.source_id_or_name = new_label
        record.errata_applied = True
        record.errata_source = append_field(record.errata_source, erratum.target)
        record.errata_version_or_date = append_field(record.errata_version_or_date, erratum.version)
        summary = f"Relabelled {old_label} as {new_label}."
        record.change_summary = append_field(record.change_summary, summary)
        by_key[new_key] = record
        return record, summary

    return None


def append_field(existing: str, addition: str) -> str:
    addition = addition.strip()
    if not addition:
        return existing
    if not existing:
        return addition
    parts = [part.strip() for part in existing.split(" | ")]
    if addition in parts:
        return existing
    return existing + " | " + addition


def log_row(erratum: Erratum, record: Record | None, status: str, applied: bool, summary: str) -> dict[str, object]:
    return {
        "target": erratum.target,
        "codex_group": erratum.codex_group,
        "page_ref": erratum.page_ref,
        "errata_version_or_date": erratum.version,
        "action": erratum.action,
        "matched_stable_key": record.stable_key if record else "",
        "matched_source_id_or_name": record.source_id_or_name if record else "",
        "status": status,
        "errata_applied": applied,
        "change_summary": summary,
        "errata_text": erratum.body,
    }


def validate(records: list[Record], errata_log: list[dict[str, object]]) -> list[str]:
    issues: list[str] = []
    keys = [record.stable_key for record in records]
    duplicates = sorted(key for key, count in defaultdict(int, ((key, keys.count(key)) for key in set(keys))).items() if count > 1)
    if duplicates:
        issues.append(f"Duplicate stable_key values: {', '.join(duplicates[:20])}")

    sections_with_records = {record.section for record in records}
    for toc in TOC:
        if toc.section not in sections_with_records:
            issues.append(f"No records found for TOC section {toc.section}.")

    outcome_ids = [int(record.source_id_or_name) for record in records if record.section == "OUTCOMES" and record.source_id_or_name.isdigit()]
    if outcome_ids != sorted(outcome_ids):
        issues.append("Outcome entries are not sorted by numeric ID.")

    unresolved = [row for row in errata_log if row["status"] == "manual_review"]
    if unresolved:
        issues.append(f"{len(unresolved)} Codex errata items require manual review; see data/errata_application_log.csv.")

    return issues


def record_sort_key(record: Record) -> tuple[int, int, str]:
    order = TOC_ORDER.get((record.section, record.subsection))
    if order is None:
        order = min((idx for (section, _), idx in TOC_ORDER.items() if section == record.section), default=999)
    if record.section == "OUTCOMES" and record.source_id_or_name.isdigit():
        return (order, int(record.source_id_or_name), record.stable_key)
    return (order, record.page_start, record.stable_key)


def write_outputs(records: list[Record], errata_log: list[dict[str, object]], out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = [record.as_dict() for record in records]

    json_path = out_dir / "stonesaga_codex_corrected.json"
    json_path.write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    csv_path = out_dir / "stonesaga_codex_corrected.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)

    xlsx_path = out_dir / "stonesaga_codex_corrected.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        pd.DataFrame(rows, columns=FIELDNAMES).to_excel(writer, index=False, sheet_name="codex")
        pd.DataFrame(errata_log).to_excel(writer, index=False, sheet_name="errata_log")

    log_path = out_dir / "errata_application_log.csv"
    with log_path.open("w", newline="", encoding="utf-8") as handle:
        fieldnames = list(errata_log[0].keys()) if errata_log else [
            "target",
            "codex_group",
            "page_ref",
            "errata_version_or_date",
            "action",
            "matched_stable_key",
            "matched_source_id_or_name",
            "status",
            "errata_applied",
            "change_summary",
            "errata_text",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(errata_log)


def write_readme(out_dir: Path) -> None:
    readme = """# Stonesaga Codex Corrected Dataset

This directory is generated by `scripts/build_codex_dataset.py` from:

- `assets/Stonesaga Codex v1.pdf`
- `assets/Stonesaga Errata & FAQ v1.3.1.docx`

## Extraction

The Codex PDF is laid out as two printed pages per PDF spread. The script crops
each spread into left and right halves, detects the printed page number in each
half, assigns that printed page to the table-of-contents ranges, and removes
repeated page headers. Outcomes are split on three-digit entry IDs and continued
entries are merged. Omen and recipe sections are split on visible entry codes
such as `FP12` or `Y2222`. Later reference/chart sections are preserved as
section-level records when the PDF text layer does not expose reliable row
boundaries.

## Errata Matching

The DOCX parser reads only the `Codex Corrections & Clarifications` section.
Rulebook and component errata are ignored unless they appear inside that Codex
section. Errata headings are mapped to stable keys by target type: outcomes,
omens, recipes, epoch/reference sections, special regions, and hex keys.

Clear generic instructions are applied directly to `corrected_text`, including
quoted text replacements, relabeling, replacing a whole entry when the errata
says it should read a full block, prepending/appending explicit added sentences,
and creating entries that the errata says should exist. Ambiguous instructions
are left unchanged and recorded in `errata_application_log.csv` with
`manual_review` status.

## Known Limitations

The PDF text layer omits or blanks some icon glyphs. Existing textual icon
placeholders from the errata DOCX are preserved, but missing PDF-only icons may
remain as spacing artifacts. Some dense chart/reference sections are emitted as
section-level records rather than one row per visual table row because the PDF
does not expose stable table boundaries. Manual review rows in the errata log
should be checked against the source PDF before publishing a fully authoritative
dataset.

## Rerun

From the repository root:

```powershell
python scripts/build_codex_dataset.py
```

The script regenerates:

- `data/stonesaga_codex_corrected.json`
- `data/stonesaga_codex_corrected.csv`
- `data/stonesaga_codex_corrected.xlsx`
- `data/errata_application_log.csv`
"""
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "README.md").write_text(readme, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pdf", type=Path, default=DEFAULT_PDF)
    parser.add_argument("--errata", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_DATA_DIR)
    args = parser.parse_args()

    records = extract_records(args.pdf)
    records.sort(key=record_sort_key)
    errata = parse_codex_errata(args.errata)
    errata_log = apply_errata(records, errata)
    records.sort(key=record_sort_key)

    issues = validate(records, errata_log)
    write_outputs(records, errata_log, args.out_dir)
    write_readme(args.out_dir)

    print(f"records={len(records)}")
    print(f"errata_items={len(errata)}")
    print(f"errata_log_rows={len(errata_log)}")
    for issue in issues:
        print(f"VALIDATION: {issue}")

    return 0 if not [issue for issue in issues if not issue.startswith("0 ")] else 0


if __name__ == "__main__":
    raise SystemExit(main())
